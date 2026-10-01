"""
Reads internal docs from a local folder: markdown, plain text, Word (.docx),
and basic text-based PDFs. This covers the common case of "I exported our
Confluence space / Notion pages / docs to a folder" before you wire up a
live connector.
"""

from pathlib import Path

from .base import BaseSource, Document


def _read_docx(path: Path) -> str:
    import docx  # python-docx

    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_pdf(path: Path) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


class LocalDocsSource(BaseSource):
    name = "local_docs"

    def is_configured(self) -> bool:
        path = self.settings.get("path")
        return bool(path) and Path(path).exists()

    def fetch(self) -> list[Document]:
        root = Path(self.settings["path"]).resolve()
        extensions = set(self.settings.get("extensions", [".md", ".txt", ".docx", ".pdf"]))

        docs: list[Document] = []

        for file_path in root.rglob("*"):
            if not file_path.is_file() or file_path.suffix.lower() not in extensions:
                continue

            try:
                suffix = file_path.suffix.lower()
                if suffix in (".md", ".txt"):
                    text = file_path.read_text(encoding="utf-8", errors="ignore")
                elif suffix == ".docx":
                    text = _read_docx(file_path)
                elif suffix == ".pdf":
                    text = _read_pdf(file_path)
                else:
                    continue
            except Exception:
                # Don't let one bad doc kill the whole indexing run.
                continue

            if not text.strip():
                continue

            rel_path = str(file_path.relative_to(root)).replace("\\", "/")
            docs.append(
                Document(
                    text=text,
                    source=f"docs/{rel_path}",
                    source_id=f"docs::{rel_path}",
                    metadata={"type": "doc"},
                )
            )

        return docs
